"""週次レポート（Markdown / CSV / Excel）。氏名・連絡先は載せない（応募者IDのみ）。"""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

from . import metrics as M
from .constants import BUSINESS_LABEL
from .jobs import readiness
from .store import Store, load_media_master, load_schedule, load_settings


def week_start(d: dt.date) -> dt.date:
    return d - dt.timedelta(days=d.weekday())


def snapshot(store: Store, today: dt.date) -> dict:
    """レポートとダッシュボードが共通で使う集計結果。"""
    jobs, settings = store.jobs(), load_settings()
    media_master, schedule = load_media_master(), load_schedule()
    apps, history, mw = store.applicants(), store.history(), store.media_weekly()
    hidx = M.history_index(history)
    jrows = M.job_rows(jobs, apps, hidx, mw)
    ws = week_start(today)
    return {
        "today": today,
        "week_start": ws,
        "settings": settings,
        "jobs": jobs,
        "apps": apps,
        "total": M.funnel(apps, hidx),
        "pace": M.pace(settings, today),
        "week": M.week_counts(apps, history, ws),
        "prev_week": M.week_counts(apps, history, ws - dt.timedelta(days=7)),
        "business": M.business_progress(apps, hidx, settings),
        "job_rows": jrows,
        "media": M.media_summary(apps, hidx, mw, media_master),
        "media_master": media_master,
        "findings": M.diagnose(jrows, apps, hidx, settings),
        "todo": M.todo(apps, today, settings),
        "phase": M.current_phase(schedule, today),
        "next_phases": M.phases_between(schedule, ws + dt.timedelta(days=7), ws + dt.timedelta(days=13)),
        "readiness": {j["job_id"]: readiness(j) for j in jobs},
    }


def _n(v) -> str:
    return "—" if v is None else f"{v:,}" if isinstance(v, int) else str(v)


def issues_and_actions(s: dict) -> tuple[list[str], list[str]]:
    """数字から課題と翌週の行動を作る（提案。実行は人が判断）。"""
    issues, actions = [], []
    t, p = s["total"], s["pace"]
    if t["応募"] < p["応募"]:
        issues.append(f"累計応募 {t['応募']}件が計画ペース {p['応募']}件を下回っている")
    for b in s["business"]:
        bf = b["_funnel"]
        if b["不足"] and bf["応募"] < b["目標"] * 2:
            name = BUSINESS_LABEL[b["事業"]]
            issues.append(f"{name}: 入社確定 {b['入社確定']}/{b['目標']}人、応募 {bf['応募']}件で母数が不足")
            actions.append(f"{name}の求人タイトルと給与・勤務時間の記載を確認し、応募対象に合わせたタイトル改善案を社長に提出する")
    not_live = [j for j in s["job_rows"] if j["公開状況"] != "公開確認済み" and not j["求人ID"].startswith("FX")]
    if not_live:
        issues.append(f"公開を確認できていない求人が {len(not_live)}件（" + "、".join(j["求人ID"] for j in not_live) + "）")
        blocked = [j["求人ID"] for j in not_live if s["readiness"][j["求人ID"]].startswith("公開不可")]
        if blocked:
            actions.append("要確認が残る求人（" + "、".join(blocked) + "）の条件を事業責任者に確認する（docs/確認事項リスト.md）")
        actions.append("社長が公開した求人は、ログアウトした状態の求職者画面で表示を確認し `publish` コマンドで「公開確認済み」を記録する")
    late = s["todo"]["初回連絡が遅れている"]
    if late:
        issues.append(f"初回連絡が{s['settings']['first_contact_sla_days']}日以上遅れている応募が {len(late)}件")
        actions.append("連絡待ちの応募者に当日中に連絡する（" + "、".join(a["応募者ID"] for a in late) + "）")
    if s["todo"]["日程調整待ち"]:
        actions.append("日程調整中の応募者に面接候補日を3つ提示する（" +
                       "、".join(a["応募者ID"] for a in s["todo"]["日程調整待ち"]) + "）")
    if s["todo"]["内定回答待ち"]:
        actions.append("内定回答待ちの方に条件の不明点がないか確認の連絡をする（" +
                       "、".join(a["応募者ID"] for a in s["todo"]["内定回答待ち"]) + "）")
    for f in s["findings"][:5]:
        issues.append(f"{f['求人ID']} {f['症状']}（{f['根拠']}）")
        actions.append(f"{f['求人ID']}: {f['改善案'][0]}")
    for ph in s["next_phases"]:
        actions.append(f"［計画］{ph['フェーズ']}: {ph['主な作業']}")
    if not issues:
        issues.append("大きな課題なし（計画どおり）")
    return issues, list(dict.fromkeys(actions))


def build_markdown(s: dict) -> str:
    ws, t, p, w, pw = s["week_start"], s["total"], s["pace"], s["week"], s["prev_week"]
    we = ws + dt.timedelta(days=6)
    issues, actions = issues_and_actions(s)
    L = [
        f"# 週次採用レポート（{ws:%Y/%m/%d}〜{we:%m/%d}）",
        "",
        f"作成日: {s['today']:%Y-%m-%d}　／　採用期限まで残り {p['days_left']}日　／　"
        f"現在のフェーズ: {s['phase']['フェーズ'] if s['phase'] else '—'}",
        "",
        "> 氏名・連絡先は載せていません（応募者IDのみ）。数字は採用管理データ（private/）から自動集計。",
        "",
        "## 1. サマリー",
        "",
        "| 指標 | 今週 | 先週 | 累計 | 計画ペース（今日時点） | 最終目標 |",
        "|---|---:|---:|---:|---:|---:|",
        f"| 応募 | {w['応募']} | {pw['応募']} | {t['応募']} | {p['応募']} | {s['settings']['funnel_targets']['応募']} |",
        f"| 面接実施 | {w['面接実施']} | {pw['面接実施']} | {t['面接実施']} | {p['面接実施']} | {s['settings']['funnel_targets']['面接実施']} |",
        f"| 内定 | {w['内定']} | {pw['内定']} | {t['内定']} | {p['内定']} | {s['settings']['funnel_targets']['内定']} |",
        f"| 入社確定 | {w['入社確定']} | {pw['入社確定']} | {t['入社確定']} | {p['入社確定']} | {s['settings']['funnel_targets']['入社確定']} |",
        "",
        f"- 面接予定（現在）: {t['面接予定']}件（うち来週 {w['面接予定_翌週']}件）　選考中: {t['選考中']}件　"
        f"不採用: {t['不採用']}件　辞退: {t['辞退']}件",
        f"- 入社済み: {t['入社済み']}人（在籍中 {t['定着中']}人／退職 {t['退職']}人）",
        f"- 応募→面接 {M.pct(M.rate(t['面接実施'], t['応募']))}　面接→内定 {M.pct(M.rate(t['内定'], t['面接実施']))}　"
        f"内定→入社確定 {M.pct(M.rate(t['入社確定'], t['内定']))}",
        "",
        "## 2. 事業別の採用状況",
        "",
        "| 事業 | 目標 | 入社確定 | 不足 | 応募 | 面接実施 | 内定 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for b in s["business"]:
        f = b["_funnel"]
        L.append(f"| {BUSINESS_LABEL[b['事業']]} | {b['目標']} | {b['入社確定']} | {b['不足']} | {f['応募']} | {f['面接実施']} | {f['内定']} |")
    tg = sum(b["目標"] for b in s["business"])
    L.append(f"| **合計** | **{tg}** | **{t['入社確定']}** | **{max(tg - t['入社確定'], 0)}** | {t['応募']} | {t['面接実施']} | {t['内定']} |")
    L += ["", "## 3. 求人別の応募状況", "",
          "| 枠 | 求人ID | 職種 | 公開状況 | 公開準備 | 表示 | クリック | 応募 | 面接 | 内定 | 確定 | 目標 | 不足 |",
          "|---:|---|---|---|---|---:|---:|---:|---:|---:|---:|---:|---:|"]
    for j in s["job_rows"]:
        L.append(f"| {j['枠']} | {j['求人ID']} | {j['職種']} | {j['公開状況']} | {s['readiness'][j['求人ID']]} | "
                 f"{_n(j['表示数'])} | {_n(j['クリック数'])} | {j['応募']} | {j['面接実施']} | {j['内定']} | "
                 f"{j['入社確定']} | {j['目標']} | {j['不足']} |")
    L += ["", "## 4. 媒体別の成果", "",
          "| 媒体 | 連携状況 | 表示 | クリック | 応募 | 面接 | 入社確定 | 広告費 | 採用単価 |",
          "|---|---|---:|---:|---:|---:|---:|---:|---:|"]
    for m in s["media"]:
        cph = "—" if m["採用単価"] is None else f"{m['採用単価']:,.0f}円"
        L.append(f"| {m['媒体']} | {m['連携状況']} | {_n(m['表示数'])} | {_n(m['クリック数'])} | {m['応募']} | "
                 f"{m['面接実施']} | {m['入社確定']} | {m['広告費']:,}円 | {cph} |")
    L += ["", "「連携状況: 未確認」の媒体は、求職者側の画面で掲載を確認できていません（掲載済みとして扱わない）。",
          "", "## 5. 課題", ""]
    L += [f"- {x}" for x in issues]
    L += ["", "## 6. 改善案（求人の変更・公開・広告費は社長の承認後に実施）", ""]
    if s["findings"]:
        for f in s["findings"]:
            L.append(f"- **{f['求人ID']} {f['職種']} — {f['症状']}**（{f['根拠']}）")
            L += [f"  - {a}" for a in f["改善案"]]
    else:
        L.append("- 公開確認済みの求人の数字がまだないため、診断は来週以降")
    L += ["", "## 7. 翌週の行動計画", ""]
    L += [f"{i}. {a}" for i, a in enumerate(actions, 1)] or ["1. —"]
    L += ["", "## 8. 対応待ち（件数）", ""]
    for k, v in s["todo"].items():
        L.append(f"- {k}: {len(v)}件" + (f"（{'、'.join(a['応募者ID'] for a in v)}）" if v else ""))
    L.append("")
    return "\n".join(L)


def write_csv(s: dict, path: Path) -> None:
    rows = [["区分", "名前", "目標", "応募", "面接実施", "内定", "入社確定", "不足", "表示数", "クリック数", "広告費"]]
    t = s["total"]
    rows.append(["全体", "合計", s["settings"]["funnel_targets"]["入社確定"], t["応募"], t["面接実施"], t["内定"],
                 t["入社確定"], max(s["settings"]["funnel_targets"]["入社確定"] - t["入社確定"], 0), "", "", ""])
    for b in s["business"]:
        f = b["_funnel"]
        rows.append(["事業", BUSINESS_LABEL[b["事業"]], b["目標"], f["応募"], f["面接実施"], f["内定"], b["入社確定"], b["不足"], "", "", ""])
    for j in s["job_rows"]:
        rows.append(["求人", f"{j['求人ID']} {j['職種']}", j["目標"], j["応募"], j["面接実施"], j["内定"], j["入社確定"],
                     j["不足"], "" if j["表示数"] is None else j["表示数"], "" if j["クリック数"] is None else j["クリック数"], j["広告費"]])
    for m in s["media"]:
        rows.append(["媒体", m["媒体"], "", m["応募"], m["面接実施"], "", m["入社確定"], "",
                     "" if m["表示数"] is None else m["表示数"], "" if m["クリック数"] is None else m["クリック数"], m["広告費"]])
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        csv.writer(f).writerows(rows)


def write_xlsx(s: dict, path: Path) -> bool:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        return False
    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F3A5F")
    head_font = Font(bold=True, color="FFFFFF", name="Meiryo")
    body_font = Font(name="Meiryo")

    def sheet(ws, title, header, rows):
        ws.title = title
        ws.append(header)
        for r in rows:
            ws.append(r)
        for c in ws[1]:
            c.fill, c.font, c.alignment = head_fill, head_font, Alignment(horizontal="center")
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
        for col in ws.columns:
            width = max(len(str(c.value or "")) for c in col)
            ws.column_dimensions[col[0].column_letter].width = min(max(8, width * 2), 60)
        ws.freeze_panes = "A2"

    t, p = s["total"], s["pace"]
    ft = s["settings"]["funnel_targets"]
    sheet(wb.active, "サマリー", ["指標", "今週", "累計", "計画ペース", "最終目標"],
          [[k, s["week"][k], t[k], p[k], ft[k]] for k in ("応募", "面接実施", "内定", "入社確定")])
    sheet(wb.create_sheet(), "事業別", ["事業", "目標", "入社確定", "不足", "応募", "面接実施", "内定"],
          [[BUSINESS_LABEL[b["事業"]], b["目標"], b["入社確定"], b["不足"], b["_funnel"]["応募"],
            b["_funnel"]["面接実施"], b["_funnel"]["内定"]] for b in s["business"]])
    sheet(wb.create_sheet(), "求人別", ["枠", "求人ID", "職種", "公開状況", "表示", "クリック", "応募", "面接", "内定", "確定", "目標", "不足"],
          [[j["枠"], j["求人ID"], j["職種"], j["公開状況"], j["表示数"], j["クリック数"], j["応募"], j["面接実施"],
            j["内定"], j["入社確定"], j["目標"], j["不足"]] for j in s["job_rows"]])
    sheet(wb.create_sheet(), "媒体別", ["媒体", "連携状況", "表示", "クリック", "応募", "面接", "入社確定", "広告費", "採用単価"],
          [[m["媒体"], m["連携状況"], m["表示数"], m["クリック数"], m["応募"], m["面接実施"], m["入社確定"],
            m["広告費"], m["採用単価"]] for m in s["media"]])
    issues, actions = issues_and_actions(s)
    sheet(wb.create_sheet(), "課題と行動", ["区分", "内容"],
          [["課題", x] for x in issues] + [["翌週の行動", x] for x in actions])
    wb.save(path)
    return True


def weekly(store: Store, today: dt.date, out_dir: Path, formats: list[str]) -> list[Path]:
    s = snapshot(store, today)
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = f"weekly_{s['week_start']:%Y-%m-%d}"
    written = []
    if "md" in formats:
        p = out_dir / f"{stem}.md"
        p.write_text(build_markdown(s), encoding="utf-8")
        written.append(p)
    if "csv" in formats:
        p = out_dir / f"{stem}.csv"
        write_csv(s, p)
        written.append(p)
    if "xlsx" in formats:
        p = out_dir / f"{stem}.xlsx"
        if write_xlsx(s, p):
            written.append(p)
    return written
